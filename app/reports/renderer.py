"""Report renderers for different output formats."""

import io
from abc import ABC, abstractmethod
from datetime import datetime
from typing import Any

import polars as pl


class BaseRenderer(ABC):
    """Abstract base class for report renderers."""

    @abstractmethod
    def render(self, report_data: dict[str, Any]) -> Any:
        """Render report to target format."""
        pass


class MarkdownRenderer(BaseRenderer):
    """Renderer for Markdown output."""

    def render(self, report_data: dict[str, Any]) -> str:
        """Render report as Markdown."""
        lines = []
        lines.append(f"# {report_data.get('report_name', 'Report')}")
        lines.append(f"\n*Generated: {report_data.get('generated_at', datetime.utcnow().isoformat())}*")
        lines.append(f"\n**Type:** {report_data.get('report_type', 'unknown')}")

        metrics = report_data.get("metrics", {})

        if report_data.get("report_type") == "dashboard":
            lines.append("\n## Dashboard\n")
            dashboard = report_data.get("dashboard", {})
            for metric_name, data in dashboard.items():
                lines.append(f"### {metric_name}")
                lines.append(f"- **Value:** {data.get('formatted_value', data.get('value'))}")
                trend = data.get("trend")
                if trend is not None:
                    trend_str = f"{trend:+.1f}%"
                    lines.append(f"- **Trend:** {trend_str}")
                sparkline = data.get("sparkline")
                if sparkline:
                    lines.append(f"- **Sparkline:** `{' '.join(f'{v:.1f}' for v in sparkline)}`")
                lines.append("")

        elif metrics:
            lines.append("\n## Metrics\n")
            lines.append("| Metric | Value |")
            lines.append("|--------|-------|")

            for metric_name, metric_data in metrics.items():
                if "value" in metric_data and metric_data["value"] is not None:
                    value = f"{metric_data['value']:.2f}"
                elif "data" in metric_data:
                    value = f"{metric_data['total_rows']} rows"
                else:
                    value = "N/A"
                lines.append(f"| {metric_name} | {value} |")

            for metric_name, metric_data in metrics.items():
                if "data" in metric_data and metric_data.get("data"):
                    lines.append(f"\n### {metric_name} Details\n")
                    df = pl.DataFrame(metric_data["data"])
                    lines.append(self._dataframe_to_markdown(df))

        return "\n".join(lines)

    def _dataframe_to_markdown(self, df: pl.DataFrame) -> str:
        """Convert DataFrame to Markdown table."""
        if df.is_empty():
            return "*No data*"

        lines = []
        lines.append("| " + " | ".join(df.columns) + " |")
        lines.append("|" + "|".join(["---"] * len(df.columns)) + "|")

        for row in df.iter_rows():
            lines.append("| " + " | ".join(str(v) for v in row) + " |")

        return "\n".join(lines)


class ExcelRenderer(BaseRenderer):
    """Renderer for Excel output."""

    def render(self, report_data: dict[str, Any]) -> bytes:
        """Render report as Excel file."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = report_data.get("report_name", "Report")[:31]

        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        row = 1
        ws.cell(row=row, column=1, value="Report: " + report_data.get("report_name", ""))
        row += 1
        ws.cell(row=row, column=1, value="Generated: " + report_data.get("generated_at", ""))
        row += 1
        ws.cell(row=row, column=1, value="Type: " + report_data.get("report_type", ""))
        row += 2

        metrics = report_data.get("metrics", {})

        if metrics:
            ws.cell(row=row, column=1, value="Metrics")
            row += 1

            for col, header in enumerate(["Metric", "Value"], start=1):
                cell = ws.cell(row=row, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font

            row += 1

            for metric_name, metric_data in metrics.items():
                if "value" in metric_data and metric_data["value"] is not None:
                    value = metric_data["value"]
                elif "data" in metric_data:
                    value = f"{metric_data['total_rows']} rows"
                else:
                    value = "N/A"

                ws.cell(row=row, column=1, value=metric_name)
                ws.cell(row=row, column=2, value=value)
                row += 1

            if report_data.get("report_type") == "dashboard":
                row += 1
                ws.cell(row=row, column=1, value="Dashboard Data")
                row += 1

                dashboard = report_data.get("dashboard", {})
                for col, header in enumerate(["Metric", "Value", "Trend", "Formatted Value"], start=1):
                    cell = ws.cell(row=row, column=col, value=header)
                    cell.fill = header_fill
                    cell.font = header_font

                row += 1

                for metric_name, data in dashboard.items():
                    ws.cell(row=row, column=1, value=metric_name)
                    ws.cell(row=row, column=2, value=data.get("value"))
                    ws.cell(row=row, column=3, value=data.get("trend"))
                    ws.cell(row=row, column=4, value=data.get("formatted_value"))
                    row += 1

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()


class ImageRenderer(BaseRenderer):
    """Renderer for image/chart output."""

    def render(self, report_data: dict[str, Any]) -> bytes:
        """Render report as PNG image with charts."""
        import plotly.graph_objects as go
        from plotly.subplots import make_subplots

        dashboard = report_data.get("dashboard", {})
        if not dashboard:
            return self._render_empty()

        num_metrics = len(dashboard)
        fig = make_subplots(
            rows=num_metrics, cols=1,
            subplot_titles=list(dashboard.keys()),
            vertical_spacing=0.1,
        )

        for idx, (metric_name, data) in enumerate(dashboard.items(), start=1):
            value = data.get("value", 0) or 0
            trend = data.get("trend")
            sparkline = data.get("sparkline")

            if sparkline:
                fig.add_trace(
                    go.Scatter(
                        y=sparkline,
                        mode="lines+markers",
                        name=metric_name,
                        line=dict(color=self._get_color(idx)),
                    ),
                    row=idx, col=1
                )
            else:
                fig.add_trace(
                    go.Indicator(
                        value=value,
                        mode="number",
                        name=metric_name,
                    ),
                    row=idx, col=1
                )

        fig.update_layout(
            height=300 * num_metrics,
            showlegend=False,
            title_text=report_data.get("report_name", "Dashboard"),
        )

        return fig.to_image(format="png")

    def _get_color(self, idx: int) -> str:
        """Get color for chart based on index."""
        colors = ["#4472C4", "#ED7D31", "#A5A5A5", "#FFC000", "#70AD47"]
        return colors[(idx - 1) % len(colors)]

    def _render_empty(self) -> bytes:
        """Render empty chart."""
        import plotly.graph_objects as go

        fig = go.Figure()
        fig.add_annotation(text="No data available", showarrow=False)
        return fig.to_image(format="png")
